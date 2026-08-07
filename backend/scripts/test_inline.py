"""Phase 4 原位转译测试：md 的 ![图] 位置 + pptx 页内图片原位插入。

注意：截图是【合成数据】（Pillow 画的模拟 VPN 报错窗口）。

验证点：
1. md 原位：图引用保留 + 转译块（###）紧随其后，前后文不被拆散
   ——关键断言：splitter 切出的【同一个子块】里同时含 前文/图转译/后文
2. 外链图片（https://）原样保留不转译
3. pptx 原位：图片转译块插回所属页 section（页码锚点）
4. 入库 + 检索闭环 + 幂等
"""
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.db import SessionLocal
from app.models import KbDocument
from app.rag.loader import load_doc
from app.rag.retriever import retrieve
from app.rag.splitter import split_all
from app.rag.store import ChromaStore
from app.rag.sync import run_sync

TMP_KB = Path(tempfile.mkdtemp(prefix="kb_inline_"))
TEST_KB_NAME = "inline_test"
TEST_STORE = ChromaStore(persist_dir=str(TMP_KB / "chroma"))


def reset_ledger() -> None:
    with SessionLocal() as db:
        db.query(KbDocument).filter_by(kb_name=TEST_KB_NAME).delete()
        db.commit()


def make_screenshot(path: Path) -> None:
    img = Image.new("RGB", (640, 300), (40, 44, 54))
    d = ImageDraw.Draw(img)
    d.rectangle([60, 40, 580, 260], fill=(24, 28, 36), outline=(120, 128, 140))
    d.text((90, 70), "VPN Client  v4.2.1", fill=(200, 210, 225))
    d.text((90, 110), "Error 800 - Connection failed", fill=(255, 90, 90))
    d.text((90, 140), "Certificate is expired", fill=(225, 230, 240))
    d.text((90, 170), "错误代码: 800", fill=(225, 230, 240))
    img.save(path, format="PNG")


MD_BODY = """---
scenario: vpn
status: active
valid_to: "2099-12-31"
---

# VPN 故障排查

## 错误码 800 场景

如下图所示，这是 VPN 客户端报错弹窗：

![VPN 报错截图](error-800.png)

图中显示证书已过期，需要申请续期。

另见外部文档：![外部图](https://example.com/other.png)
"""


def main() -> None:
    reset_ledger()

    # ---- 0) 准备：md + 同目录截图 + 带图 pptx ----
    vpn_dir = TMP_KB / "vpn"
    vpn_dir.mkdir(parents=True, exist_ok=True)
    make_screenshot(vpn_dir / "error-800.png")
    (vpn_dir / "guide.md").write_text(MD_BODY, encoding="utf-8")

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "VPN 培训"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "错误码 800 处理流程"
    s2.placeholders[1].text = "步骤 1：查询证书状态"
    s2.shapes.add_picture(str(vpn_dir / "error-800.png"), Inches(1), Inches(3))
    prs.save(vpn_dir / "with-image.pptx")
    print("[准备] guide.md（含 ![图]）+ error-800.png + with-image.pptx（第 2 页含图）")

    # ---- 1) md 原位转译 ----
    doc = load_doc(vpn_dir / "guide.md", TMP_KB)
    assert doc is not None
    c = doc.content
    assert "![VPN 报错截图](error-800.png)" in c, "图引用必须保留"
    assert "https://example.com/other.png" in c, "外链图片必须原样保留"
    assert "### 截图转译：error-800.png" in c, "应插入 ### 转译块"
    assert "800" in c.split("### 截图转译")[1][:300], "转译块应含错误码 800"
    # 顺序：前文 → 图引用 → 转译块 → 后文
    order = [c.index(s) for s in ("如下图所示", "![VPN 报错截图]", "### 截图转译", "图中显示")]
    assert order == sorted(order), f"上下文顺序被拆散: {order}"
    print("[md-原位] 顺序正确（前文→引用→转译→后文），转译块含错误码 800")

    # ---- 2) 关键断言：同一个子块里同时有 前文/转译/后文（上下文连贯） ----
    children = split_all([doc])[0].children
    assert len(children) == 1, f"### 不应切新章节，子块数应仍为 1: {len(children)}"
    blk = children[0].text
    assert all(s in blk for s in ("如下图所示", "### 截图转译", "图中显示")), \
        "前文/转译/后文必须在同一子块（上下文连贯）"
    print("[splitter] 子块数不变（1 个），同一子块内含 前文+转译+后文 ✅")

    # ---- 3) pptx 页码锚点：转译块在第二页 section 内 ----
    from app.rag.loader import load_pptx_doc
    pdoc = load_pptx_doc(vpn_dir / "with-image.pptx", TMP_KB)
    assert pdoc is not None and "### 截图转译" in pdoc.content
    second_page_start = pdoc.content.rindex("## 错误码 800 处理流程")
    assert pdoc.content.index("### 截图转译", second_page_start) > second_page_start, \
        "pptx 图片转译块应位于第二页 section 之后"
    after_second = pdoc.content[second_page_start + len("## 错误码 800 处理流程"):]
    new_chapters = [ln for ln in after_second.splitlines() if ln.startswith("## ")]
    assert not new_chapters, f"第二页之后不应再有新章节（图是 ### 内联块）: {new_chapters}"
    print("[pptx-原位] 图片转译块插入第二页 section 内（页码锚点生效）")

    # ---- 4) 入库 + 检索 + 幂等 ----
    report1 = run_sync(kb_root=TMP_KB, store=TEST_STORE, kb_name=TEST_KB_NAME)
    print(f"[sync#1] {report1.summary}")
    assert len(report1.added) == 3, f"应新增 3 篇（md+png+pptx）: {report1.summary}"
    hits = retrieve("VPN 客户端 错误码 800 证书过期", scenario="vpn", top_k=3,
                    store=TEST_STORE)
    top = hits[0] if hits else None
    assert top is not None, "检索应命中"
    print(f"[检索] Top1: {top.id}，命中来源: {top.metadata.get('doc_id')}（chunk 含上下文）")
    assert "截图转译" in top.parent_text or "截图转译" in top.text, "命中块应含转译内容"
    report2 = run_sync(kb_root=TMP_KB, store=TEST_STORE, kb_name=TEST_KB_NAME)
    assert len(report2.skipped) == 3, "第二次应全跳过"

    print("\n=== Phase 4 原位转译（md ![图] + pptx 页码锚点）：全部通过 ===")


if __name__ == "__main__":
    main()
